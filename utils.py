from datetime import datetime
from pathlib import Path


def export_chat(messages, format='txt'):
    base_dir = Path(__file__).parent
    export_dir = base_dir / 'export'
    export_dir.mkdir(exist_ok=True)
    logo_path = base_dir / 'static' / 'logo-dark.png'
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if format == 'txt':
        content = ""
        for msg in messages:
            content += f"\n[{msg['role'].upper()}]\n{msg['content']}\n"
            if msg.get('context'):
                content += "\nReferenced Documents:\n"
                for ctx in msg['context']:
                    content += f"- [{ctx['id']}] {ctx['title']} (Relevance: {float(ctx['occurrence_weight']):.2f})\n"
            content += "\n---\n"
        filename = export_dir / f"chat_export_{timestamp}.txt"
        filename.write_text(content, encoding='utf-8')

    elif format == 'docx':
        from docx import Document
        from docx.shared import Inches
        doc = Document()

        if logo_path.exists():
            doc.add_picture(str(logo_path), width=Inches(1.5))
            doc.add_paragraph()

        for msg in messages:
            doc.add_heading(f"{msg['role'].upper()}", level=2)
            doc.add_paragraph(msg['content'])
            if msg.get('context'):
                doc.add_paragraph("Referenced Documents:")
                for ctx in msg['context']:
                    doc.add_paragraph(
                        f"[{ctx['id']}] {ctx['title']} (Relevance: {float(ctx['occurrence_weight']):.2f})",
                        style='List Bullet'
                    )
            doc.add_paragraph("---")
        filename = export_dir / f"chat_export_{timestamp}.docx"
        doc.save(str(filename))

    elif format == 'pdf':
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()

        if logo_path.exists():
            pdf.image(str(logo_path), x=170, y=10, w=30)
            pdf.ln(35)

        pdf.set_font("Arial", size=12)
        for msg in messages:
            pdf.cell(200, 10, txt=f"{msg['role'].upper()}", ln=True)
            pdf.multi_cell(200, 10, txt=msg['content'])
            if msg.get('context'):
                pdf.cell(200, 10, txt="Referenced Documents:", ln=True)
                for ctx in msg['context']:
                    pdf.cell(200, 10,
                        txt=f"[{ctx['id']}] {ctx['title']} (Relevance: {float(ctx['occurrence_weight']):.2f})",
                        ln=True
                    )
            pdf.cell(200, 10, txt="---", ln=True)
        filename = export_dir / f"chat_export_{timestamp}.pdf"
        pdf.output(str(filename))

    return str(filename)

